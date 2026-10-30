import socket
import struct
import json
import os
import io
import hashlib
import tkinter as tk

from tkinter import filedialog

from Crypto.Cipher import DES
from Crypto.Util.Padding import pad, unpad

from pypdf import PdfReader
from docx import Document


# ============================================================
# SETTINGS
# ============================================================

HOST = "0.0.0.0"
PORT = 5000

KEY = b"8bytekey"


# ============================================================
# DES ENCRYPTION
# ============================================================

def encrypt_data(data):

    cipher = DES.new(
        KEY,
        DES.MODE_CBC
    )

    ciphertext = cipher.encrypt(
        pad(data, DES.block_size)
    )

    return cipher.iv + ciphertext


# ============================================================
# DES DECRYPTION
# ============================================================

def decrypt_data(data):

    iv = data[:8]
    ciphertext = data[8:]

    cipher = DES.new(
        KEY,
        DES.MODE_CBC,
        iv
    )

    plaintext = unpad(
        cipher.decrypt(ciphertext),
        DES.block_size
    )

    return plaintext


# ============================================================
# HASH
# ============================================================

def get_hash(data):

    return hashlib.sha256(data).hexdigest()


# ============================================================
# SEND HEADER
# ============================================================

def send_header(sock, data):

    message = json.dumps(
        data
    ).encode("utf-8")

    sock.sendall(
        struct.pack(
            "!I",
            len(message)
        )
    )

    sock.sendall(message)


# ============================================================
# RECEIVE HEADER
# ============================================================

def receive_header(sock):

    size_data = receive_exact(
        sock,
        4
    )

    if not size_data:
        return None

    size = struct.unpack(
        "!I",
        size_data
    )[0]

    data = receive_exact(
        sock,
        size
    )

    if not data:
        return None

    return json.loads(
        data.decode("utf-8")
    )


# ============================================================
# RECEIVE EXACT DATA
# ============================================================

def receive_exact(sock, size):

    data = b""

    while len(data) < size:

        packet = sock.recv(
            min(
                65536,
                size - len(data)
            )
        )

        if not packet:
            return None

        data += packet

    return data


# ============================================================
# SELECT FILE FROM ANY LOCATION
# ============================================================

def select_file():

    root = tk.Tk()
    root.withdraw()

    filepath = filedialog.askopenfilename(
        title="Select File",
        filetypes=[
            ("All Files", "*.*"),
            ("Text Files", "*.txt"),
            ("PDF Files", "*.pdf"),
            ("Word Files", "*.docx"),
            ("Images", "*.jpg;*.jpeg;*.png"),
            ("Binary Files", "*.bin")
        ]
    )

    root.destroy()

    return filepath


# ============================================================
# DISPLAY CONTENT
# ============================================================

def display_content(filename, data):

    extension = os.path.splitext(
        filename
    )[1].lower()

    print("\n")
    print("=" * 70)
    print("PLAINTEXT / DECRYPTED CONTENT")
    print("=" * 70)

    print(
        "File:",
        filename
    )

    print(
        "Size:",
        len(data),
        "bytes"
    )

    # --------------------------------------------------------
    # TEXT
    # --------------------------------------------------------

    if extension in [
        ".txt",
        ".csv",
        ".py",
        ".json",
        ".html",
        ".xml",
        ".log"
    ]:

        print("\nTEXT CONTENT")
        print("-" * 70)

        text = data.decode(
            "utf-8",
            errors="replace"
        )

        print(text)

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    elif extension == ".pdf":

        print("\nPDF TEXT CONTENT")
        print("-" * 70)

        try:

            pdf = PdfReader(
                io.BytesIO(data)
            )

            print(
                "Pages:",
                len(pdf.pages)
            )

            for number, page in enumerate(
                pdf.pages,
                start=1
            ):

                print(
                    f"\n--- PAGE {number} ---"
                )

                text = page.extract_text()

                if text:

                    print(text)

                else:

                    print(
                        "[No selectable text]"
                    )

        except Exception as e:

            print(
                "PDF extraction error:",
                e
            )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    elif extension == ".docx":

        print("\nWORD DOCUMENT CONTENT")
        print("-" * 70)

        try:

            document = Document(
                io.BytesIO(data)
            )

            for paragraph in document.paragraphs:

                if paragraph.text.strip():

                    print(
                        paragraph.text
                    )

            for table in document.tables:

                print("\n[TABLE]")

                for row in table.rows:

                    values = []

                    for cell in row.cells:

                        values.append(
                            cell.text
                        )

                    print(
                        " | ".join(values)
                    )

        except Exception as e:

            print(
                "DOCX extraction error:",
                e
            )

    # --------------------------------------------------------
    # BINARY
    # --------------------------------------------------------

    else:

        print("\nBINARY FILE")
        print("-" * 70)

        print(
            "This file is binary data."
        )

        print(
            "Hexadecimal preview:"
        )

        print(
            data[:200].hex()
        )

    print("=" * 70)


# ============================================================
# SERVER SEND FILE
# ============================================================

def send_file(conn):

    filepath = select_file()

    if not filepath:

        print(
            "\nNo file selected."
        )

        return

    filename = os.path.basename(
        filepath
    )

    # Read original file
    with open(
        filepath,
        "rb"
    ) as file:

        plaintext = file.read()

    print("\n")
    print("=" * 70)
    print("SERVER → CLIENT")
    print("=" * 70)

    print(
        "Selected:",
        filepath
    )

    print(
        "Plaintext size:",
        len(plaintext),
        "bytes"
    )

    print(
        "SHA-256:",
        get_hash(plaintext)
    )

    # Encrypt ONLY IN MEMORY
    ciphertext = encrypt_data(
        plaintext
    )

    print(
        "\n✓ DES encryption successful"
    )

    print(
        "Ciphertext size:",
        len(ciphertext),
        "bytes"
    )

    print(
        "\nCiphertext preview:"
    )

    print(
        ciphertext[:100].hex()
    )

    # Send header
    send_header(
        conn,
        {
            "operation": "SEND_FILE",
            "filename": filename,
            "filesize": len(ciphertext)
        }
    )

    # Send encrypted bytes
    conn.sendall(
        ciphertext
    )

    print(
        "\n✓ Ciphertext sent to client"
    )

    print(
        "✓ Encrypted file was NOT saved"
    )


# ============================================================
# SERVER RECEIVE FILE
# ============================================================

def receive_file(conn):

    print("\n")
    print("=" * 70)
    print("SERVER WAITING FOR CLIENT FILE")
    print("=" * 70)

    # Tell client that server is ready
    send_header(
        conn,
        {
            "operation": "READY"
        }
    )

    header = receive_header(
        conn
    )

    if header is None:

        print(
            "Client disconnected."
        )

        return

    if header["operation"] != "SEND_FILE":

        print(
            "Unexpected response."
        )

        return

    filename = header[
        "filename"
    ]

    filesize = header[
        "filesize"
    ]

    ciphertext = receive_exact(
        conn,
        filesize
    )

    if ciphertext is None:

        print(
            "File transfer failed."
        )

        return

    print("\n")
    print("=" * 70)
    print("CLIENT → SERVER")
    print("=" * 70)

    print(
        "File:",
        filename
    )

    print(
        "Ciphertext size:",
        len(ciphertext),
        "bytes"
    )

    print(
        "\nCiphertext preview:"
    )

    print(
        ciphertext[:100].hex()
    )

    # Decrypt ONLY IN MEMORY
    try:

        plaintext = decrypt_data(
            ciphertext
        )

    except Exception as e:

        print(
            "\nDES decryption error:",
            e
        )

        return

    print(
        "\n✓ DES decryption successful"
    )

    # Display actual content
    display_content(
        filename,
        plaintext
    )

    print(
        "\n✓ Encrypted file was NOT saved"
    )

    print(
        "✓ Decrypted file was NOT saved"
    )


# ============================================================
# SERVER
# ============================================================

server = socket.socket(
    socket.AF_INET,
    socket.SOCK_STREAM
)

server.setsockopt(
    socket.SOL_SOCKET,
    socket.SO_REUSEADDR,
    1
)

server.bind(
    (
        HOST,
        PORT
    )
)

server.listen(1)

print("\n")
print("=" * 70)
print("                 DES SERVER")
print("              PyCryptodome")
print("=" * 70)

print(
    "\nWaiting for client..."
)

conn, address = server.accept()

print(
    "\n✓ Client connected:",
    address
)


# ============================================================
# SERVER LOOP
# ============================================================

while True:

    try:

        header = receive_header(
            conn
        )

        if header is None:

            print(
                "\nClient disconnected."
            )

            break

        operation = header[
            "operation"
        ]

        # ----------------------------------------------------
        # CLIENT SENDS FILE
        # ----------------------------------------------------

        if operation == "SEND_FILE":

            filename = header[
                "filename"
            ]

            filesize = header[
                "filesize"
            ]

            ciphertext = receive_exact(
                conn,
                filesize
            )

            if ciphertext is None:

                break

            print("\n")
            print("=" * 70)
            print("CLIENT → SERVER")
            print("=" * 70)

            print(
                "File:",
                filename
            )

            print(
                "Ciphertext size:",
                len(ciphertext),
                "bytes"
            )

            print(
                "\nCiphertext preview:"
            )

            print(
                ciphertext[:100].hex()
            )

            try:

                plaintext = decrypt_data(
                    ciphertext
                )

            except Exception as e:

                print(
                    "\nDES decryption error:",
                    e
                )

                continue

            print(
                "\n✓ DES decryption successful"
            )

            display_content(
                filename,
                plaintext
            )

            print(
                "\n✓ Encrypted file was NOT saved"
            )

            print(
                "✓ Decrypted file was NOT saved"
            )

        # ----------------------------------------------------
        # CLIENT REQUESTS SERVER FILE
        # ----------------------------------------------------

        elif operation == "REQUEST_FILE":

            send_file(
                conn
            )

        # ----------------------------------------------------
        # EXIT
        # ----------------------------------------------------

        elif operation == "EXIT":

            print(
                "\nClient exited."
            )

            break

    except Exception as e:

        print(
            "\nServer error:",
            e
        )

        break


conn.close()
server.close()

print(
    "\nServer stopped."
)