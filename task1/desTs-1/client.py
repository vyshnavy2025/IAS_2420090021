import socket
import struct
import json
import os
import io
import hashlib
import tkinter as tk

from tkinter import filedialog

from Crypto.Cipher import DES
from Crypto.Util.Padding import pad, unpad

from pypdf import PdfReader
from docx import Document


# ============================================================
# SETTINGS
# ============================================================

SERVER_IP = "127.0.0.1"
PORT = 5000

# DES requires exactly 8 bytes
KEY = b"8bytekey"


# ============================================================
# DES ENCRYPTION
# ============================================================

def encrypt_data(data):

    cipher = DES.new(
        KEY,
        DES.MODE_CBC
    )

    ciphertext = cipher.encrypt(
        pad(data, DES.block_size)
    )

    # IV + ciphertext
    return cipher.iv + ciphertext


# ============================================================
# DES DECRYPTION
# ============================================================

def decrypt_data(data):

    iv = data[:8]
    ciphertext = data[8:]

    cipher = DES.new(
        KEY,
        DES.MODE_CBC,
        iv
    )

    plaintext = unpad(
        cipher.decrypt(ciphertext),
        DES.block_size
    )

    return plaintext


# ============================================================
# HASH
# ============================================================

def get_hash(data):

    return hashlib.sha256(data).hexdigest()


# ============================================================
# SEND HEADER
# ============================================================

def send_header(sock, data):

    message = json.dumps(
        data
    ).encode("utf-8")

    sock.sendall(
        struct.pack(
            "!I",
            len(message)
        )
    )

    sock.sendall(message)


# ============================================================
# RECEIVE HEADER
# ============================================================

def receive_header(sock):

    size_data = receive_exact(
        sock,
        4
    )

    if not size_data:
        return None

    size = struct.unpack(
        "!I",
        size_data
    )[0]

    data = receive_exact(
        sock,
        size
    )

    if not data:
        return None

    return json.loads(
        data.decode("utf-8")
    )


# ============================================================
# RECEIVE EXACT DATA
# ============================================================

def receive_exact(sock, size):

    data = b""

    while len(data) < size:

        packet = sock.recv(
            min(
                65536,
                size - len(data)
            )
        )

        if not packet:
            return None

        data += packet

    return data


# ============================================================
# SELECT ANY FILE
# ============================================================

def select_file():

    root = tk.Tk()
    root.withdraw()

    filepath = filedialog.askopenfilename(
        title="Select File",
        filetypes=[
            ("All Files", "*.*"),
            ("Text Files", "*.txt"),
            ("PDF Files", "*.pdf"),
            ("Word Files", "*.docx"),
            ("Images", "*.jpg;*.jpeg;*.png"),
            ("Binary Files", "*.bin")
        ]
    )

    root.destroy()

    return filepath


# ============================================================
# DISPLAY FILE CONTENT
# ============================================================

def display_content(filename, data):

    extension = os.path.splitext(
        filename
    )[1].lower()

    print("\n")
    print("=" * 70)
    print("PLAINTEXT / DECRYPTED CONTENT")
    print("=" * 70)

    print("File:", filename)
    print("Size:", len(data), "bytes")

    # --------------------------------------------------------
    # TEXT FILE
    # --------------------------------------------------------

    if extension in [
        ".txt",
        ".csv",
        ".py",
        ".json",
        ".html",
        ".xml",
        ".log"
    ]:

        print("\nTEXT CONTENT")
        print("-" * 70)

        text = data.decode(
            "utf-8",
            errors="replace"
        )

        print(text)

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    elif extension == ".pdf":

        print("\nPDF TEXT CONTENT")
        print("-" * 70)

        try:

            pdf = PdfReader(
                io.BytesIO(data)
            )

            print(
                "Pages:",
                len(pdf.pages)
            )

            for number, page in enumerate(
                pdf.pages,
                start=1
            ):

                print(
                    f"\n--- PAGE {number} ---"
                )

                text = page.extract_text()

                if text:

                    print(text)

                else:

                    print(
                        "[No selectable text]"
                    )

        except Exception as e:

            print(
                "PDF extraction error:",
                e
            )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    elif extension == ".docx":

        print("\nWORD DOCUMENT CONTENT")
        print("-" * 70)

        try:

            document = Document(
                io.BytesIO(data)
            )

            # Paragraphs
            for paragraph in document.paragraphs:

                if paragraph.text.strip():

                    print(
                        paragraph.text
                    )

            # Tables
            for table in document.tables:

                print("\n[TABLE]")

                for row in table.rows:

                    values = []

                    for cell in row.cells:

                        values.append(
                            cell.text
                        )

                    print(
                        " | ".join(values)
                    )

        except Exception as e:

            print(
                "DOCX extraction error:",
                e
            )

    # --------------------------------------------------------
    # BINARY FILE
    # --------------------------------------------------------

    else:

        print("\nBINARY FILE")
        print("-" * 70)

        print(
            "This file is binary data."
        )

        print(
            "Hexadecimal preview:"
        )

        print(
            data[:200].hex()
        )

    print("=" * 70)


# ============================================================
# CLIENT SEND FILE
# ============================================================

def send_file(sock):

    filepath = select_file()

    if not filepath:

        print(
            "\nNo file selected."
        )

        return

    filename = os.path.basename(
        filepath
    )

    # Read original file
    with open(
        filepath,
        "rb"
    ) as file:

        plaintext = file.read()

    print("\n")
    print("=" * 70)
    print("CLIENT → SERVER")
    print("=" * 70)

    print(
        "Selected:",
        filepath
    )

    print(
        "Plaintext size:",
        len(plaintext),
        "bytes"
    )

    original_hash = get_hash(
        plaintext
    )

    print(
        "SHA-256:",
        original_hash
    )

    # Encrypt ONLY IN MEMORY
    ciphertext = encrypt_data(
        plaintext
    )

    print(
        "\n✓ DES encryption successful"
    )

    print(
        "Ciphertext size:",
        len(ciphertext),
        "bytes"
    )

    print(
        "\nCiphertext preview:"
    )

    print(
        ciphertext[:100].hex()
    )

    # Send header
    send_header(
        sock,
        {
            "operation": "SEND_FILE",
            "filename": filename,
            "filesize": len(ciphertext)
        }
    )

    # Send encrypted bytes
    sock.sendall(
        ciphertext
    )

    print(
        "\n✓ Ciphertext sent to server"
    )

    print(
        "✓ Encrypted file was NOT saved"
    )

    print(
        "Waiting for server..."
    )


# ============================================================
# CLIENT RECEIVE FILE
# ============================================================

def receive_file(sock):

    print("\n")
    print("=" * 70)
    print("CLIENT REQUESTING FILE FROM SERVER")
    print("=" * 70)

    # Ask server to send a file
    send_header(
        sock,
        {
            "operation": "REQUEST_FILE"
        }
    )

    response = receive_header(
        sock
    )

    if response is None:

        print(
            "Server disconnected."
        )

        return

    if response["operation"] == "NO_FILE":

        print(
            "\nServer did not send a file."
        )

        return

    filename = response[
        "filename"
    ]

    filesize = response[
        "filesize"
    ]

    print(
        "\nReceiving:",
        filename
    )

    print(
        "Ciphertext size:",
        filesize,
        "bytes"
    )

    # Receive ciphertext ONLY IN MEMORY
    ciphertext = receive_exact(
        sock,
        filesize
    )

    if ciphertext is None:

        print(
            "File transfer failed."
        )

        return

    print(
        "\n✓ Ciphertext received"
    )

    print(
        "Ciphertext preview:"
    )

    print(
        ciphertext[:100].hex()
    )

    # Decrypt ONLY IN MEMORY
    try:

        plaintext = decrypt_data(
            ciphertext
        )

    except Exception as e:

        print(
            "\nDES decryption error:",
            e
        )

        return

    print(
        "\n✓ DES decryption successful"
    )

    # Display actual content
    display_content(
        filename,
        plaintext
    )

    print(
        "\n✓ File was NOT saved"
    )


# ============================================================
# CLIENT
# ============================================================

print("\n")
print("=" * 70)
print("                 DES CLIENT")
print("              PyCryptodome")
print("=" * 70)

client = socket.socket(
    socket.AF_INET,
    socket.SOCK_STREAM
)

try:

    client.connect(
        (
            SERVER_IP,
            PORT
        )
    )

    print(
        "\n✓ Connected to server"
    )

except Exception as e:

    print(
        "\nConnection failed:",
        e
    )

    exit()


# ============================================================
# CLIENT MENU
# ============================================================

while True:

    print("\n")
    print("=" * 50)
    print("CLIENT MENU")
    print("=" * 50)

    print(
        "1. Send File"
    )

    print(
        "2. Receive File"
    )

    print(
        "3. Exit"
    )

    choice = input(
        "\nEnter choice: "
    )

    if choice == "1":

        send_file(
            client
        )

    elif choice == "2":

        receive_file(
            client
        )

    elif choice == "3":

        send_header(
            client,
            {
                "operation": "EXIT"
            }
        )

        print(
            "\nClient stopped."
        )

        break

    else:

        print(
            "\nInvalid choice."
        )


client.close()